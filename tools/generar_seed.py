"""
Genera el seed de SIGMA Mantención a partir de las planillas verificadas del caso A1 · Minera Cerro Sauce.

Cada fila de detención de las planillas se convierte en una orden de trabajo correctiva, con las mismas
horas, equipo, turno, contratista, causa y descripción: así lo que Copilot responda desde SIGMA calza
fila por fila con lo que la clase muestra desde la planilla de SharePoint.

Se agregan además:
  - la OT «pendiente de repuesto» del revestimiento del chute de CV-03 (05-06-2026), el hilo del caso;
  - la intervención menor del 18-06-2026 en CV-03 (40 min, turno noche) que la bitácora registra y la
    planilla no — por eso no genera registro de detención;
  - preventivas programadas por área y mes (no generan detención), para que el tablero se vea vivo.

Uso:
  python tools/generar_seed.py "<carpeta A1 - Cerro Sauce>"
Escribe sigma/seed/ordenes.json y sigma/seed/costos.json. Se versionan: el despliegue no necesita las planillas.
"""
import json
import pathlib
import random
import sys

import openpyxl

RAIZ = pathlib.Path(__file__).resolve().parents[1]
SALIDA = RAIZ / "sigma" / "seed"

EQUIPOS_PREVENTIVA = {
    "Chancado": ["Chancador primario CH-01", "Correa CV-02", "Correa CV-03", "Harnero HR-01", "Alimentador AL-02"],
    "Molienda": ["Molino SAG ML-01", "Molino de bolas ML-02", "Bomba de pulpa BP-04"],
    "Planta Concentradora": ["Celda de flotación FL-04", "Espesador ES-01", "Bomba BP-07"],
    "Mina": ["Pala PC-01", "Camión CAEX-11", "Perforadora PF-02"],
    "Despacho": ["Romana RM-01", "Cargador frontal CF-03"],
}
CONTRATISTA_PREVENTIVA = {"Chancado": "Mantención Andes S.A.", "Molienda": "Mantención Andes S.A.",
                          "Planta Concentradora": "Servicios Eléctricos Sur", "Mina": "Propio", "Despacho": "Propio"}
TAREAS_PREVENTIVA = ["Inspección programada y lubricación", "Cambio de aceite y filtros",
                     "Medición de desgaste de revestimientos", "Revisión de alineación y tensado",
                     "Termografía de tableros y motores", "Mantención de 500 horas"]


def leer(planilla: pathlib.Path) -> list[dict]:
    ws = openpyxl.load_workbook(planilla, data_only=True)["Detenciones"]
    filas = list(ws.iter_rows(values_only=True))[1:]
    return [{"fecha": r[1].date().isoformat(), "mes": r[2], "area": r[3], "equipo": r[4], "turno": r[5],
             "contratista": r[6], "causa": r[7], "horas": float(r[8]), "descripcion": r[9], "registrado_por": r[10]}
            for r in filas if r[0]]


def main(a1: pathlib.Path) -> None:
    det = leer(a1 / "CerroSauce-Detenciones-2026.xlsx")
    vistas = {(d["fecha"], d["equipo"], d["horas"]) for d in det}
    for d in leer(a1 / "CerroSauce-Detenciones-Julio.xlsx"):
        if (d["fecha"], d["equipo"], d["horas"]) not in vistas:
            det.append(d)
    det.sort(key=lambda d: d["fecha"])

    ordenes = []
    for d in det:
        ordenes.append({
            "fecha": d["fecha"] + ("T02:00:00" if d["turno"] == "Noche" else "T10:00:00"),
            "area": d["area"], "equipo": d["equipo"], "tipo": "Correctiva", "turno": d["turno"],
            "causa": d["causa"], "descripcion": d["descripcion"], "horas": d["horas"],
            "contratista": d["contratista"], "estado": "Cerrada",
            "costo_repuestos_clp": 0, "detencion_registrada": True, "origen": "planilla",
            "registrado_por": d["registrado_por"], "nota": "",
        })

    # Los dos hilos del caso que la planilla no tiene
    ordenes.append({
        "fecha": "2026-06-05T11:00:00", "area": "Chancado", "equipo": "Correa CV-03", "tipo": "Correctiva",
        "turno": "Día", "causa": "Falta de repuesto",
        "descripcion": "Reemplazo de revestimiento del chute de descarga de CV-03 para eliminar atollos recurrentes.",
        "horas": 0.0, "contratista": "Mantención Andes S.A.", "estado": "Pendiente de repuesto",
        "costo_repuestos_clp": 4800000, "detencion_registrada": False, "origen": "sigma", "registrado_por": "R. Aliaga",
        "nota": "Repuesto solicitado el 05-06-2026. Sin fecha de llegada confirmada por el proveedor.",
    })
    ordenes.append({
        "fecha": "2026-06-18T02:05:00", "area": "Chancado", "equipo": "Correa CV-03", "tipo": "Correctiva menor",
        "turno": "Noche", "causa": "Error operacional",
        "descripcion": "Atollo en chute de descarga. Limpieza manual, 40 minutos. Registrado por turno noche como intervención menor.",
        "horas": 0.67, "contratista": "Mantención Andes S.A.", "estado": "Cerrada",
        "costo_repuestos_clp": 0, "detencion_registrada": False, "origen": "sigma", "registrado_por": "Turno noche",
        "nota": "Clasificada como intervención menor: no genera registro de detención en la planilla de Operaciones.",
    })

    # Preventivas deterministas, una por área y mes, marzo–julio 2026
    rnd = random.Random(2026)
    for mes in ["2026-03", "2026-04", "2026-05", "2026-06", "2026-07"]:
        for area, equipos in EQUIPOS_PREVENTIVA.items():
            dia = rnd.randint(2, 27)
            if area == "Chancado" and mes in ("2026-06", "2026-07"):
                continue  # los meses plantados del caso se dejan limpios: solo lo que la planilla y la bitácora dicen
            ordenes.append({
                "fecha": f"{mes}-{dia:02d}T09:00:00", "area": area, "equipo": rnd.choice(equipos),
                "tipo": "Preventiva", "turno": "Día", "causa": "Programada",
                "descripcion": rnd.choice(TAREAS_PREVENTIVA) + ".", "horas": float(rnd.choice([4, 5, 6, 8, 10])),
                "contratista": CONTRATISTA_PREVENTIVA[area], "estado": "Cerrada",
                "costo_repuestos_clp": rnd.choice([0, 850000, 1200000, 2300000, 3900000]),
                "detencion_registrada": False, "origen": "sigma", "registrado_por": "Planificación", "nota": "",
            })

    # Costo de repuestos plausible para las correctivas por falla mecánica / falta de repuesto
    for o in ordenes:
        if o["tipo"] == "Correctiva" and o["costo_repuestos_clp"] == 0 and o["causa"] in ("Falla mecánica", "Falta de repuesto"):
            o["costo_repuestos_clp"] = int(o["horas"] * rnd.choice([350000, 600000, 900000]) // 10000 * 10000)

    ordenes.sort(key=lambda o: o["fecha"])
    for i, o in enumerate(ordenes, start=1):
        o["id"] = f"OT-{o['fecha'][:4]}-{i:04d}"

    SALIDA.mkdir(parents=True, exist_ok=True)
    (SALIDA / "ordenes.json").write_text(json.dumps(ordenes, ensure_ascii=False, indent=1), encoding="utf-8")

    ws = openpyxl.load_workbook(a1 / "Costo-Hora-Indisponibilidad-2026.xlsx", data_only=True)["Costo hora 2026"]
    costos = [{"area": r[0], "costo_hora_clp": int(r[1]), "base": r[2], "vigencia": r[3]}
              for r in list(ws.iter_rows(values_only=True))[1:] if r[0]]
    (SALIDA / "costos.json").write_text(json.dumps(costos, ensure_ascii=False, indent=1), encoding="utf-8")

    ch6 = [o for o in ordenes if o["area"] == "Chancado" and o["fecha"].startswith("2026-06") and o["detencion_registrada"]]
    ch7 = [o for o in ordenes if o["area"] == "Chancado" and o["fecha"].startswith("2026-07") and o["detencion_registrada"]]
    print(f"{len(ordenes)} OT · {len(det)} desde planillas · Chancado jun {sum(o['horas'] for o in ch6)} h ({len(ch6)}) · "
          f"jul {sum(o['horas'] for o in ch7)} h ({len(ch7)}) · pendiente: "
          f"{next(o['id'] for o in ordenes if o['estado'] == 'Pendiente de repuesto')} · menor 18-06: "
          f"{next(o['id'] for o in ordenes if o['tipo'] == 'Correctiva menor')}")


if __name__ == "__main__":
    main(pathlib.Path(sys.argv[1]))
